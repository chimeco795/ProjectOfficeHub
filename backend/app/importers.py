"""Deterministic proposals. Document content is data and is never executed."""
import re
import unicodedata
from io import BytesIO
from datetime import date, datetime
from zipfile import ZipFile, BadZipFile
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

SECTIONS = ['general','avance','logros','problemas','actividades','riesgos','hitos','bloqueos','dependencias','proximos_pasos','decisiones','alertas','acciones']
ALIASES = {
 'description': ['actividad','tarea','riesgo','descripcion','nombre','hito','pendiente anterior'],
 'code': ['id','codigo'], 'owner': ['responsable','owner','asignado a'],
 'status': ['estatus','estado','estado actual'], 'start_date':['inicio','fecha inicio'],
 'end_date':['fin','fecha fin','fecha compromiso','fecha plan'],
 'probability':['probabilidad'], 'impact':['impacto'], 'mitigation':['mitigacion','accion mitigacion'],
 'consequence':['consecuencia'], 'justification':['justificacion'],
 'milestone':['hito'], 'dependencies':['dependencias notas','dependencias','predecesoras'],
 'category':['fase','categoria'], 'weight':['ponderacion'], 'critical':['ruta critica','critica'],
 'planned':['avance planeado','planeado','% estimado'], 'actual':['avance real','real','% alcanzado']
}

def norm(value):
    return re.sub(r'[^a-z0-9%]+', ' ', ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').lower()) if not unicodedata.combining(c))).strip()

def serial(value):
    return value.isoformat() if isinstance(value, (datetime,date)) else value

def validate_zip(data):
    try:
        with ZipFile(BytesIO(data)) as z:
            if len(z.infolist()) > 10000 or sum(i.file_size for i in z.infolist()) > 100_000_000:
                raise ValueError('Documento demasiado grande al descomprimir (máximo 100 MB).')
    except BadZipFile:
        raise ValueError('El archivo no es un documento Office válido.')

def proposal(section, location, original, current, doubtful=False):
    return dict(section=section, location=location, original=original, current=current,
                review='dudoso' if doubtful else 'pendiente')

def word(data):
    doc = Document(BytesIO(data))
    records, section, heading, pending = [], 'general', 'Documento', None
    def flush():
        nonlocal pending
        if pending:
            records.append(proposal(section, pending['location'], {'text': '\n'.join(pending['text']), 'heading':heading}, {'description':'\n'.join(pending['text'])}))
            pending = None
    markers = [('logros','logros'),('problemas','problemas'),('estado de los pendientes','actividades'),('proximos pasos','proximos_pasos'),('evaluacion','general'),('riesgos','riesgos'),('bloqueos','bloqueos'),('decisiones','decisiones'),('conclusion','general')]
    for index, el in enumerate(doc.element.body):
        location = f'Bloque {index+1}'
        if el.tag.endswith('}tbl'):
            flush()
            table = Table(el, doc)
            rows = [[c.text for c in r.cells] for r in table.rows]
            for n,row in enumerate(rows[1:],2):
                records.append(proposal(section, f'{heading} / {location}, fila {n}', {'headers':rows[0],'cells':row}, {'description':row[0], 'status':row[1] if len(row)>1 else ''}))
            continue
        if not el.tag.endswith('}p'): continue
        p = Paragraph(el,doc)
        text = p.text.strip()
        if not text: continue
        normalized = norm(text)
        marker = next((s for prefix,s in markers if normalized.startswith(prefix)),None)
        if marker and (len(text)<100 or normalized.startswith('conclusion')):
            flush(); section=marker; heading=text
            if ':' not in text: continue
        if re.match(r'^\d+[.)]\s',text):
            flush(); pending={'location':f'{heading} / {location}', 'text':[text]}; continue
        if pending:
            pending['text'].append(text); continue
        if normalized in ['criticos','construccion']:
            heading=text; continue
        metrics = {}
        for field,label in [('planned',r'planeado'),('actual',r'real'),('source_variation',r'variaci[oó]n')]:
            match=re.search(label+r'\s*:\s*(-?\d+(?:[.,]\d+)?)\s*%',text,re.I)
            if match: metrics[field]=float(match[1].replace(',','.'))
        records.append(proposal('avance' if metrics else section, f'{heading} / {location}', {'text':text}, metrics or {'description':text}))
    flush()
    return records, ['Extracción narrativa por encabezados y numeración. Responsables y fechas implícitos requieren revisión.']

def excel(data, mapping):
    wb=load_workbook(BytesIO(data),data_only=False)
    cached=load_workbook(BytesIO(data),data_only=True)
    records,warnings=[],[]
    for sheet in wb:
        if sheet.max_row * sheet.max_column > 250000:
            raise ValueError(f'Hoja {sheet.title}: supera el límite de 250.000 celdas.')
        rows=list(sheet.iter_rows())
        config=mapping.get(sheet.title,{})
        # Horizontal timeline detected by labels, independent of sheet name.
        labels={norm(c.value):(r,c.column) for r,row in enumerate(rows) for c in row[:3] if isinstance(c.value,str)}
        if 'fecha corte' in labels and '% estimado' in labels and not config:
            dr,dc=labels['fecha corte']; pr,_=labels['% estimado']; ar=labels.get('% alcanzado',(None,None))[0]
            for col in range(dc,sheet.max_column):
                dt=rows[dr][col].value
                if dt is None: continue
                original={}
                current={'date':serial(dt)}
                for key,ri in [('date',dr),('planned',pr),('actual',ar)]:
                    if ri is None:continue
                    cell=rows[ri][col]; value=cached[sheet.title][cell.coordinate].value
                    original[cell.coordinate]={'value':serial(cell.value),'cached':serial(value),'format':cell.number_format}
                    if key!='date': current[key]=round(value*100,6) if isinstance(value,(float,int)) else None
                records.append(proposal('avance',f'{sheet.title} / columna {col+1}',original,current))
            warnings.append(f'{sheet.title}: proyecciones conservadas como datos; no se crean cortes históricos automáticamente.')
            continue
        def detect(row):
            result={}
            for cell in row:
                for key,aliases in ALIASES.items():
                    if norm(cell.value) in aliases: result.setdefault(key,cell.column-1)
            return result
        if config:
            hr=int(config.get('header_row',1))-1
            if hr<0 or hr>=len(rows):raise ValueError('Fila de encabezados fuera del rango.')
            columns=config.get('columns',{})
            header={key:next((c.column-1 for c in rows[hr] if str(c.value)==str(label)), -1) for key,label in columns.items()}
            if any(i<0 for i in header.values()):raise ValueError(f'{sheet.title}: un encabezado configurado no existe.')
        else:
            candidates=[(len(detect(row)),idx,detect(row)) for idx,row in enumerate(rows[:30])]
            _,hr,header=max(candidates,default=(0,0,{}),key=lambda x:x[0])
        section=config.get('section') or ('riesgos' if 'probability' in header else 'actividades')
        if section not in SECTIONS:raise ValueError('Sección de mapeo desconocida.')
        if len(header)<2:
            warnings.append(f'{sheet.title}: estructura no reconocida; filas conservadas sin clasificación. Configure el mapeo en otra importación.')
            hr=-1;section='general'
        for row in rows[hr+1:]:
            if not any(c.value is not None for c in row):continue
            original={c.coordinate:{'header':serial(rows[hr][c.column-1].value) if hr>=0 else None,'value':serial(c.value),'cached':serial(cached[sheet.title][c.coordinate].value),'format':c.number_format} for c in row if c.value is not None}
            current={key:serial(cached[sheet.title][row[i].coordinate].value) for key,i in header.items() if row[i].value is not None}
            for key in ['planned','actual']:
                if key in current and isinstance(current[key],(int,float)) and '%' in row[header[key]].number_format:current[key]=round(current[key]*100,6)
            row_section=section
            if norm(current.get('category'))=='master plan':row_section='general'
            if not current:current={'description':' | '.join(str(c.value) for c in row if c.value is not None)}
            records.append(proposal(row_section,f'{sheet.title} / fila {row[0].row}',original,current,len(header)<2))
        unknown=[str(c.value) for c in rows[hr] if c.value is not None and c.column-1 not in header.values()] if hr>=0 else []
        if unknown:warnings.append(f'{sheet.title}: columnas sin mapear conservadas en originales: '+', '.join(unknown))
    wb.close();cached.close()
    return records,warnings

def parse(data,extension,mapping):
    validate_zip(data)
    return word(data) if extension=='.docx' else excel(data,mapping)
