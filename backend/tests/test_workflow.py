from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
import pytest
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook
from app import db
from app.main import app
from app.importers import parse

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(db,'DATA',tmp_path)
    with TestClient(app) as client:yield client

def setup(client,name='Proyecto independiente'):
    p=client.post('/api/projects',json={'name':name}).json()
    c=client.post(f"/api/projects/{p['id']}/cuts",json={'report_date':'2026-09-25'}).json()
    return p,c

def doc_bytes():
    d=Document();d.add_paragraph('Logros principales');d.add_paragraph('1. Liberación');d.add_paragraph('Pedro confirmó la liberación')
    d.add_paragraph('Evaluación ejecutiva');d.add_paragraph('Avance Planeado: 48%\nAvance Real: 45%')
    b=BytesIO();d.save(b);return b.getvalue()

def test_persistence_original_duplicates_and_isolation(client):
    payload=doc_bytes()
    p,c=setup(client)
    route=f"/api/cuts/{c['id']}/imports"
    assert client.post(route,files={'file':('a.docx',payload)}).status_code==201
    detail=client.get(f"/api/cuts/{c['id']}").json();row=detail['records'][0]
    current={**row['current'],'description':'Pablo confirmó la liberación'}
    value={'version':1,'section':'logros','review':'aceptado','current':current}
    assert client.patch('/api/records/'+row['id'],json=value).status_code==200
    assert client.patch('/api/records/'+row['id'],json=value).status_code==409
    assert client.post(route,files={'file':('renamed.docx',payload)}).status_code==409
    # Reopen app and a fresh connection: persisted state survives lifecycle restart.
    with TestClient(app) as fresh:
        saved=fresh.get(f"/api/cuts/{c['id']}").json()['records'][0]
        assert 'Pedro' in saved['original']['text'] and 'Pablo' in saved['current']['description']
    assert saved['modified']==1
    assert len(client.get('/api/records/'+row['id']+'/audit').json())==1
    _,other=setup(client,'Otro proyecto')
    assert not client.get(f"/api/cuts/{other['id']}").json()['records']
    assert client.post(f"/api/projects/{p['id']}/cuts",json={'report_date':'2026-09-25'}).status_code==409
    next_cut=client.post(f"/api/projects/{p['id']}/cuts",json={'report_date':'2026-10-02'}).json()
    assert not client.get(f"/api/cuts/{next_cut['id']}").json()['records']

def test_generic_excel_and_mapping():
    w=Workbook();s=w.active;s.title='Cualquier nombre';s.append(['Tarea','Responsable','FIN']);s.append(['Entrega','Equipo','2026-10-10'])
    t=w.create_sheet('Serie');t.append(['Fecha Corte','2026-09-25','2026-10-02']);t.append(['% Estimado',.48,.55]);t.append(['% Alcanzado',.45,None])
    b=BytesIO();w.save(b)
    rows,_=parse(b.getvalue(),'.xlsx',{})
    assert rows[0]['current']['owner']=='Equipo'
    assert rows[1]['current']['planned']==48
    assert rows[2]['current']['actual'] is None
    rows,_=parse(b.getvalue(),'.xlsx',{'Cualquier nombre':{'header_row':1,'section':'hitos','columns':{'description':'Tarea','owner':'Responsable'}}})
    assert rows[0]['section']=='hitos'

def test_validation_and_calculation(client):
    _,c=setup(client)
    route=f"/api/cuts/{c['id']}/imports"
    assert client.post(route,files={'file':('bad.xlsx',b'not zip')}).status_code==422
    assert client.post(route,files={'file':('file.exe',b'bad')}).status_code==415
    client.post(route,files={'file':('a.docx',doc_bytes())})
    row=client.get(f"/api/cuts/{c['id']}").json()['records'][-1]
    value={'version':1,'section':'avance','review':'aceptado','current':{'planned':48,'actual':101}}
    assert client.patch('/api/records/'+row['id'],json=value).status_code==422
    value['current']['actual']=45
    assert client.patch('/api/records/'+row['id'],json=value).status_code==200
    row=client.get(f"/api/cuts/{c['id']}").json()['records'][-1]
    assert row['generated']['variation']==-3

def test_soft_delete_and_restore(client):
    _,c=setup(client);client.post(f"/api/cuts/{c['id']}/imports",files={'file':('a.docx',doc_bytes())})
    row=client.get(f"/api/cuts/{c['id']}").json()['records'][0]
    for version,state in [(1,'eliminado'),(2,'pendiente')]:
        assert client.patch('/api/records/'+row['id'],json={'version':version,'section':row['section'],'review':state,'current':row['current']}).status_code==200
    assert len(client.get('/api/records/'+row['id']+'/audit').json())==2
